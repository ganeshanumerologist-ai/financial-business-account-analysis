"""Real .xlsx workbook tests: values, charts, sheets and audit traceability."""
from __future__ import annotations

import io
import os
import sys
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

from app import accounting as acct
from app import analytics, db, excel


class ExcelWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.conn = db.connect(":memory:")
        db.init_db(self.conn)
        self.t = db.create_tenant(self.conn, "Excel Co", "excel-co", state_code="27")
        self.tenant = self.conn.execute("SELECT * FROM tenant WHERE id=?", (self.t,)).fetchone()
        self.post("1000", "4000", 100000, "2026-04-10", "Sale")
        self.post("6000", "1000", 40000, "2026-05-10", "Rent")

    def tearDown(self):
        self.conn.close()

    def post(self, debit_code: str, credit_code: str, amount: int,
             entry_date: str, narration: str):
        return acct.post_journal(
            self.conn, self.t, entry_date, narration,
            [{"account_id": db.account_id(self.conn, self.t, debit_code),
              "debit": amount, "credit": 0},
             {"account_id": db.account_id(self.conn, self.t, credit_code),
              "debit": 0, "credit": amount}],
            "manual")

    def workbook(self, date_from="2026-04-01", date_to="2026-06-30"):
        data = excel.build_workbook(self.conn, self.tenant, date_from, date_to)
        self.assertTrue(data.startswith(b"PK"))
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            self.assertIn("[Content_Types].xml", zf.namelist())
            self.assertTrue(any(name.startswith("xl/worksheets/") for name in zf.namelist()))
        return data

    def test_workbook_has_full_sheet_pack_and_charts(self):
        data = self.workbook()
        wb = load_workbook(io.BytesIO(data), data_only=False)
        for name in ("Cover", "Read Me", "Data Dictionary", "Executive KPIs", "Monthly",
                     "Profit & Loss", "Balance Sheet", "Trial Balance", "Expense Breakdown",
                     "Top Customers", "Top Vendors", "AR Ageing", "AP Ageing", "Cash & Bank",
                     "Inventory", "GST Working", "GST Sales", "GST Purchases", "Journal",
                     "AI Insights"):
            self.assertIn(name, wb.sheetnames)
        self.assertGreaterEqual(len(wb["Monthly"]._charts), 2)
        self.assertGreaterEqual(len(wb["Expense Breakdown"]._charts), 1)
        self.assertEqual(wb["Cover"]["B2"].value, "ARTHIVIO")
        self.assertIn("posted journal", wb["Cover"]["A21"].value.lower())
        self.assertIn("No estimate or forecast", wb["Read Me"]["B11"].value)

    def test_kpis_and_statements_reconcile_to_posted_journal(self):
        wb = load_workbook(io.BytesIO(self.workbook()), data_only=False)
        kpi = wb["Executive KPIs"]
        values = {}
        paise = {}
        for row in range(5, kpi.max_row + 1):
            label = kpi.cell(row, 1).value
            if label:
                values[label] = kpi.cell(row, 2).value
                paise[label] = kpi.cell(row, 4).value
        self.assertAlmostEqual(values["Revenue"], 1000.00, places=2)
        self.assertAlmostEqual(values["Expenses"], 400.00, places=2)
        self.assertAlmostEqual(values["Net profit"], 600.00, places=2)
        self.assertEqual(paise["Revenue"], 100000)
        self.assertEqual(paise["Expenses"], 40000)
        self.assertEqual(paise["Net profit"], 60000)

        pnl = wb["Profit & Loss"]
        pnl_values = [pnl.cell(row, 4).value for row in range(5, pnl.max_row + 1)]
        self.assertIn(1000.0, pnl_values)
        self.assertIn(400.0, pnl_values)
        self.assertIn(600.0, pnl_values)

        tb = wb["Trial Balance"]
        self.assertEqual(tb.cell(tb.max_row - 2, 2).value, "TOTAL")
        self.assertEqual(tb.cell(tb.max_row - 2, 6).value, "Yes")

    def test_monthly_sheet_zero_fills_and_has_actual_totals(self):
        wb = load_workbook(io.BytesIO(self.workbook()), data_only=False)
        ws = wb["Monthly"]
        months = {}
        for row in range(5, ws.max_row + 1):
            key = ws.cell(row, 1).value
            if key and len(str(key)) == 7:
                months[key] = (ws.cell(row, 2).value, ws.cell(row, 3).value,
                               ws.cell(row, 4).value)
        self.assertIn("2026-04", months)
        self.assertIn("2026-06", months)
        self.assertEqual(months["2026-04"], (1000.0, 0.0, 1000.0))
        self.assertEqual(months["2026-05"], (0.0, 400.0, -400.0))
        self.assertEqual(months["2026-06"], (0.0, 0.0, 0.0))
        total_row = max(row for row in range(5, ws.max_row + 1)
                        if ws.cell(row, 1).value == "Total")
        self.assertAlmostEqual(ws.cell(total_row, 2).value, 1000.0, places=2)
        self.assertAlmostEqual(ws.cell(total_row, 3).value, 400.0, places=2)
        self.assertAlmostEqual(ws.cell(total_row, 4).value, 600.0, places=2)

    def test_data_dictionary_and_ai_insights_are_grounded(self):
        wb = load_workbook(io.BytesIO(self.workbook()), data_only=False)
        dictionary = wb["Data Dictionary"]
        keys = {dictionary.cell(row, 1).value for row in range(5, dictionary.max_row + 1)}
        self.assertIn("revenue_paise", keys)
        self.assertIn("trial_balance_balanced", keys)
        insights = wb["AI Insights"]
        text = " ".join(str(insights.cell(row, col).value or "")
                        for row in range(1, insights.max_row + 1) for col in range(1, 5))
        self.assertIn("Deterministic", text)
        self.assertIn("No generative AI provider", text)
        self.assertIn("net_profit_paise / revenue_paise", text)
        self.assertIn("Forecasts are excluded", text)

    def test_empty_period_is_zero_filled_not_invented(self):
        wb = load_workbook(io.BytesIO(
            self.workbook("2025-01-01", "2025-02-28")), data_only=False)
        monthly = wb["Monthly"]
        self.assertEqual(monthly.cell(5, 1).value, "2025-01")
        self.assertEqual(monthly.cell(5, 2).value, 0.0)
        self.assertEqual(monthly.cell(6, 1).value, "2025-02")
        self.assertEqual(monthly.cell(6, 4).value, 0.0)
        insights = wb["AI Insights"]
        text = " ".join(str(insights.cell(row, 2).value or "")
                        for row in range(5, insights.max_row + 1))
        self.assertIn("not generated", text.lower())

    def test_invalid_date_is_rejected_and_filename_is_safe(self):
        with self.assertRaises(analytics.AnalyticsError) as ctx:
            excel.build_workbook(self.conn, self.tenant, "not-a-date", "2026-06-30")
        self.assertEqual(ctx.exception.code, "BAD_DATE")
        self.assertEqual(excel.filename(self.tenant, "2026-06-30"),
                         "arthivio-excel-co-actuals-2026-06-30.xlsx")


if __name__ == "__main__":
    unittest.main()
