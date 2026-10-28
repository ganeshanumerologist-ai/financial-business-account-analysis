"""End-to-end HTTP tests: login, CSRF, invoice creation, report pages.

Runs the real server against a temporary database — no mocks.
"""
from __future__ import annotations

import html
import http.cookiejar
import io
import json
import os
import re
import sqlite3
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from http.server import ThreadingHTTPServer

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app import accounting as acct
from app import bank, db, security, server
from app.demo import DEMO_PASSWORD, seed_demo


class WebTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="arthivio-test-")
        cls.db_path = os.path.join(cls.tmp, "test.db")
        server.DB_PATH = cls.db_path
        server.BACKUP_DIR = os.path.join(cls.tmp, "backups")
        conn = db.connect(cls.db_path)
        db.init_db(conn)
        seed_demo(conn)
        cls.tenant_id = conn.execute("SELECT id FROM tenant WHERE slug='arthivio-demo'").fetchone()["id"]
        cls.customer_id = conn.execute("SELECT id FROM customer WHERE tenant_id=? LIMIT 1", (cls.tenant_id,)).fetchone()["id"]
        cls.item_id = conn.execute("SELECT id FROM item WHERE tenant_id=? AND track_stock=1 LIMIT 1", (cls.tenant_id,)).fetchone()["id"]
        conn.close()

        cls.httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        cls.port = cls.httpd.server_address[1]
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()

        cls.jar = http.cookiejar.CookieJar()
        cls.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cls.jar))

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()

    # ---------- helpers ----------
    def url(self, path):
        return f"http://127.0.0.1:{self.port}{path}"

    def get(self, path):
        with self.opener.open(self.url(path), timeout=15) as r:
            return r.status, r.read().decode("utf-8")

    def get_json(self, path):
        try:
            with self.opener.open(self.url(path), timeout=15) as r:
                return r.status, json.loads(r.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            try:
                return e.code, json.loads(e.read().decode("utf-8"))
            finally:
                e.close()

    def get_bytes(self, path):
        with self.opener.open(self.url(path), timeout=30) as r:
            return r.status, dict(r.headers), r.read()

    def post(self, path, data: dict):
        body = urllib.parse.urlencode(data, doseq=True).encode()
        req = urllib.request.Request(self.url(path), data=body, method="POST")
        try:
            with self.opener.open(req, timeout=15) as r:
                return r.status, r.read().decode("utf-8"), r.geturl()
        except urllib.error.HTTPError as e:
            try:
                return e.code, e.read().decode("utf-8"), e.geturl()
            finally:
                e.close()

    def post_multipart(self, path, fields: dict, files: dict):
        boundary = "----arthivio-test-" + os.urandom(8).hex()
        chunks = []
        for name, value in fields.items():
            chunks.append(
                f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"\r\n\r\n'
                f'{value}\r\n'.encode("utf-8")
            )
        for name, (filename, content, content_type) in files.items():
            chunks.append(
                f'--{boundary}\r\nContent-Disposition: form-data; name="{name}"; '
                f'filename="{filename}"\r\nContent-Type: {content_type}\r\n\r\n'.encode("utf-8")
            )
            chunks.append(content)
            chunks.append(b"\r\n")
        chunks.append(f"--{boundary}--\r\n".encode("utf-8"))
        body = b"".join(chunks)
        req = urllib.request.Request(
            self.url(path), data=body, method="POST",
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        )
        try:
            with self.opener.open(req, timeout=30) as r:
                return r.status, r.read().decode("utf-8"), r.geturl()
        except urllib.error.HTTPError as e:
            try:
                return e.code, e.read().decode("utf-8"), e.geturl()
            finally:
                e.close()

    def xlsx_bytes(self, rows) -> bytes:
        wb = Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def csrf(self, page: str) -> str:
        m = re.search(r'name="_csrf" value="([^"]+)"', page)
        self.assertIsNotNone(m, "CSRF token missing from page")
        return m.group(1)

    def login(self):
        status, body, _ = self.post("/login", {"email": "owner@arthivio.test", "password": DEMO_PASSWORD})
        self.assertEqual(status, 200)
        self.assertIn("Dashboard", body)

    # ---------- tests ----------
    def test_00_public_landing_page(self):
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        with opener.open(self.url("/"), timeout=15) as r:
            body = r.read().decode("utf-8")
            self.assertEqual(r.status, 200)
            self.assertNotIn("/login", r.geturl())
        self.assertIn("Smarter Finance.", body)
        self.assertIn("GLOBAL BUSINESS MANAGEMENT PLATFORM", body)
        self.assertIn("Trusted by Businesses Worldwide", body)
        self.assertIn("50,000+*", body)
        self.assertIn("illustrative marketing metrics", body)
        self.assertIn("Sign in", body)
        self.assertIn("synthetic data", body)
        self.assertNotIn("None", body)
        self.assertNotIn("NaN", body)

    def test_01_login_required(self):
        s, body = self.get("/dashboard-does-not-exist")
        self.assertIn("Sign in", body)

    def test_02_bad_password_rejected(self):
        status, body, _ = self.post("/login", {"email": "owner@arthivio.test", "password": "wrong-password"})
        self.assertEqual(status, 401)
        self.assertIn("Invalid email or password", body)

    def test_03_login_and_dashboard(self):
        self.login()
        s, body = self.get("/")
        self.assertEqual(s, 200)
        self.assertIn("Revenue", body)
        self.assertIn("Net profit", body)
        self.assertNotIn("None", body)
        self.assertNotIn("NaN", body)

    def test_04_reports_render_and_balance(self):
        self.login()
        s, tb = self.get("/reports/trial-balance")
        self.assertIn("Balanced", tb)
        self.assertNotIn("Out of balance", tb)
        s, pnl = self.get("/reports/pnl")
        self.assertIn("NET PROFIT", pnl)
        s, bs = self.get("/reports/balance-sheet")
        self.assertIn("A = L + E", bs)
        s, ledger = self.get("/ledger")
        self.assertIn("Closing balance", ledger)

    def test_05_invoice_creation_flows_to_reports(self):
        self.login()
        s, page = self.get("/invoices/new")
        token = self.csrf(page)
        status, body, final = self.post("/invoices", {
            "_csrf": token,
            "customer_id": self.customer_id,
            "invoice_date": "2026-10-22",
            "place_of_supply": "27",
            "rounding": "1",
            "item_id": [self.item_id],
            "qty": ["2"],
            "price": [""],
            "discount": ["0"],
        })
        self.assertEqual(status, 200)
        self.assertIn("Invoice total", body)
        self.assertRegex(final, r"/invoices/[0-9a-f]{32}$")

        s, tb = self.get("/reports/trial-balance")
        self.assertIn("Balanced", tb)
        s, pnl = self.get("/reports/pnl")
        self.assertIn("Sales Revenue", pnl)

    def test_06_csrf_required(self):
        self.login()
        status, body, _ = self.post("/invoices", {
            "customer_id": self.customer_id, "invoice_date": "2026-10-22",
            "item_id": [self.item_id], "qty": ["1"], "discount": ["0"],
        })
        self.assertEqual(status, 403)
        self.assertIn("CSRF", body)

    def test_07_customer_creation(self):
        self.login()
        s, page = self.get("/customers")
        token = self.csrf(page)
        status, _, _ = self.post("/customers", {"_csrf": token, "name": "Test Buyer Pvt Ltd",
                                                "phone": "9999999999", "state_code": "27", "gstin": ""})
        self.assertEqual(status, 200)
        s, page = self.get("/customers")
        self.assertIn("Test Buyer Pvt Ltd", page)

    def test_08_invoice_not_visible_across_tenants(self):
        self.login()
        conn = db.connect(self.db_path)
        other_tenant = db.create_tenant(conn, "Other Books", "other-books", state_code="29")
        from app import security
        uid = security.create_user(conn, "other@arthivio.test", "AnotherLongPass!2026", "Other")
        security.add_membership(conn, other_tenant, uid, "owner")
        foreign_invoice = conn.execute(
            "SELECT id FROM invoice WHERE tenant_id=? LIMIT 1", (self.tenant_id,)).fetchone()["id"]
        conn.close()

        other_jar = http.cookiejar.CookieJar()
        other = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(other_jar))
        req = urllib.request.Request(self.url("/login"),
                                     data=urllib.parse.urlencode({"email": "other@arthivio.test",
                                                                  "password": "AnotherLongPass!2026"}).encode(),
                                     method="POST")
        with other.open(req, timeout=15) as r:
            r.read()
        with other.open(self.url(f"/invoices/{foreign_invoice}"), timeout=15) as r:
            body = r.read().decode()
        self.assertIn("Invoice not found", body)


    def test_09_customer_statement(self):
        self.login()
        s, page = self.get("/customers")
        self.assertIn("Balance due", page)
        m = re.search(r'href="(/customers/[0-9a-f]{32})"', page)
        self.assertIsNotNone(m, "customer statement link missing")
        s, st = self.get(m.group(1))
        self.assertIn("Closing balance", st)
        self.assertIn("Statement", st)
        self.assertIn('href="/invoices/', st)


    def test_10_receipt_posts_and_reduces_outstanding(self):
        self.login()
        conn = db.connect(self.db_path)
        inv = conn.execute("SELECT id FROM invoice WHERE tenant_id=? ORDER BY invoice_no DESC LIMIT 1",
                           (self.tenant_id,)).fetchone()
        before = acct.invoice_outstanding(conn, self.tenant_id, inv["id"])
        conn.close()
        self.assertGreater(before, 100)

        page = self.get(f"/invoices/{inv['id']}")[1]
        token = self.csrf(page)
        status, body, _ = self.post("/receipts", {
            "_csrf": token, "invoice_id": inv["id"], "amount": "1.00", "mode": "bank",
            "payment_date": "2026-10-23", "reference": "TEST-RCPT"})
        self.assertEqual(status, 200)
        self.assertIn("Outstanding", body)

        conn = db.connect(self.db_path)
        after = acct.invoice_outstanding(conn, self.tenant_id, inv["id"])
        conn.close()
        self.assertEqual(after, before - 100)

    def test_11_bill_and_vendor_payment(self):
        self.login()
        page = self.get("/vendors")[1]
        token = self.csrf(page)
        self.post("/vendors", {"_csrf": token, "name": "Test Supplier", "phone": "",
                               "state_code": "27", "gstin": ""})
        conn = db.connect(self.db_path)
        vendor = conn.execute("SELECT id FROM vendor WHERE tenant_id=? AND name='Test Supplier'",
                              (self.tenant_id,)).fetchone()
        conn.close()
        self.assertIsNotNone(vendor)

        page = self.get("/bills/new")[1]
        token = self.csrf(page)
        status, body, final = self.post("/bills", {
            "_csrf": token, "vendor_id": vendor["id"], "bill_date": "2026-10-23",
            "vendor_bill_no": "T-1", "place_of_supply": "27",
            "item_id": [self.item_id], "qty": ["2"], "price": ["50"], "discount": ["0"]})
        self.assertEqual(status, 200)
        self.assertIn("Bill total", body)
        self.assertRegex(final, r"/bills/[0-9a-f]{32}$")

        bill_id = final.rsplit("/", 1)[-1]
        token = self.csrf(body)
        status, paid, _ = self.post("/receipts", {
            "_csrf": token, "direction": "payment", "bill_id": bill_id,
            "amount": "105.00", "mode": "bank", "payment_date": "2026-10-23"})
        self.assertEqual(status, 200)
        self.assertIn("Outstanding ₹0.00", paid)

    def test_12_credit_note_from_invoice(self):
        self.login()
        conn = db.connect(self.db_path)
        inv = conn.execute("SELECT id FROM invoice WHERE tenant_id=? ORDER BY invoice_no DESC LIMIT 1",
                           (self.tenant_id,)).fetchone()
        line = conn.execute("SELECT id FROM invoice_line WHERE invoice_id=? LIMIT 1",
                            (inv["id"],)).fetchone()
        conn.close()

        page = self.get(f"/invoices/{inv['id']}")[1]
        token = self.csrf(page)
        status, body, _ = self.post("/credit-notes", {
            "_csrf": token, "invoice_id": inv["id"], "invoice_line_id": line["id"],
            "qty": "1", "note_date": "2026-10-23", "reason": "Web test return"})
        self.assertEqual(status, 200)
        self.assertIn("Credit note", body)

        conn = db.connect(self.db_path)
        count = conn.execute("SELECT COUNT(*) c FROM credit_note WHERE tenant_id=? AND reason='Web test return'",
                             (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertEqual(count, 1)

    def test_13_csv_exports(self):
        self.login()
        for path, needle in [("/export/trial-balance.csv", "TOTAL"),
                             ("/export/gst-sales.csv", "Taxable"),
                             ("/export/invoices.csv", "Outstanding"),
                             ("/export/daybook.csv", "Voucher"),
                             ("/export/bank-statement.csv?from=2026-04-01&to=2026-10-24", "Status")]:
            status, body = self.get(path)
            self.assertEqual(status, 200, path)
            self.assertIn(needle, body, path)

    def test_14_new_pages_render(self):
        self.login()
        pages = {
            "/receipts": "Advance receipt",
            "/bills": "New bill",
            "/vendors": "Add vendor",
            "/daybook": "Day book",
            "/journal": "Post journal",
            "/reports/ar-ageing": "Receivables ageing",
            "/reports/ap-ageing": "Payables ageing",
            "/reports/gst": "Output tax",
            "/reports/cash-book": "Cash Book",
            "/reports/bank-book": "Bank Book",
            "/credit-notes": "credit note",
            "/backup": "Backup",
        }
        for path, needle in pages.items():
            status, body = self.get(path)
            self.assertEqual(status, 200, path)
            self.assertIn(needle, body, path)
            self.assertNotIn("Internal server error", body, path)

    def test_15_backup_over_http(self):
        self.login()
        page = self.get("/backup")[1]
        token = self.csrf(page)
        status, body, final = self.post("/backup", {"_csrf": token})
        self.assertEqual(status, 200)
        self.assertIn("Backup created and verified", body)
        self.assertIn("done=1", final)
        backup_dir = os.path.join(self.tmp, "backups")
        self.assertTrue(os.path.isdir(backup_dir))
        self.assertTrue([f for f in os.listdir(backup_dir) if f.endswith(".db")])


    def test_16_settings_and_print_invoice(self):
        self.login()
        page = self.get("/settings")[1]
        self.assertIn("Business identity", page)
        token = self.csrf(page)

        bad_status, bad_body, _ = self.post("/settings", {
            "_csrf": token, "name": "Arthivio Demo Trading", "gstin": "BAD",
            "state_code": "27"})
        self.assertEqual(bad_status, 400)
        self.assertIn("GSTIN must be 15 characters", bad_body)

        status, body, _ = self.post("/settings", {
            "_csrf": token, "name": "Arthivio Demo Trading",
            "gstin": "27AABCS1429B1Z1", "state_code": "27",
            "address": "12, Laxmi Industrial Estate", "city": "Mumbai",
            "pincode": "400053", "phone": "+91 22 4890 1200",
            "email": "accounts@arthivio.test", "bank_name": "HDFC Bank",
            "bank_account": "50200012345678", "bank_ifsc": "HDFC0001234",
            "invoice_terms": "Payment due within 15 days."})
        self.assertEqual(status, 200)
        self.assertIn("Settings saved", body)
        self.assertIn("27AABCS1429B1Z1", body)

        conn = db.connect(self.db_path)
        inv = conn.execute(
            "SELECT id, invoice_no FROM invoice WHERE tenant_id=? ORDER BY invoice_no LIMIT 1",
            (self.tenant_id,)).fetchone()
        conn.close()
        status, page = self.get(f"/invoices/{inv['id']}/print")
        self.assertEqual(status, 200)
        self.assertIn("Tax Invoice", page)
        self.assertIn(inv["invoice_no"], page)
        self.assertIn("27AABCS1429B1Z1", page)
        self.assertIn("Amount in words", page)
        self.assertIn("HDFC Bank", page)
        self.assertIn("Print / Save as PDF", page)

    def test_17_gstr1_csv_export(self):
        self.login()
        status, body = self.get("/export/gstr1.csv?from=2026-04-01&to=2026-10-24")
        self.assertEqual(status, 200)
        for needle in ("B2B", "B2CS", "HSN", "CN TOTAL", "TOTAL", "Working export"):
            self.assertIn(needle, body)
        self.assertIn("Rate %", body)

    def test_18_gstr1_json_export(self):
        self.login()
        status, body = self.get("/export/gstr1.json?from=2026-04-01&to=2026-10-24")
        self.assertEqual(status, 200)
        data = json.loads(body)
        self.assertEqual(data["format"], "arthivio-gstr1-working-v1")
        self.assertTrue(data["b2b"])
        self.assertTrue(data["hsn"])
        self.assertIsInstance(data["totals"]["taxable"], str)
        self.assertIn("not an uploaded or filed", data["note"])


    def test_19_debit_note_from_bill(self):
        self.login()
        conn = db.connect(self.db_path)
        bill = conn.execute(
            "SELECT b.id, b.total_paise, bl.id AS line_id FROM bill b"
            " JOIN bill_line bl ON bl.bill_id=b.id"
            " WHERE b.tenant_id=? ORDER BY b.bill_date DESC, b.bill_no DESC LIMIT 1",
            (self.tenant_id,)).fetchone()
        before = acct.bill_outstanding(conn, self.tenant_id, bill["id"])
        conn.close()

        page = self.get(f"/bills/{bill['id']}")[1]
        token = self.csrf(page)
        status, body, final = self.post("/debit-notes", {
            "_csrf": token, "bill_id": bill["id"], "bill_line_id": bill["line_id"],
            "qty": "1", "note_date": "2026-10-24", "reason": "Web test debit return"})
        self.assertEqual(status, 200)
        self.assertIn("Debit note", body)
        self.assertIn("Outstanding", body)
        self.assertRegex(final, rf"/bills/{bill['id']}$")

        conn = db.connect(self.db_path)
        note = conn.execute(
            "SELECT * FROM debit_note WHERE tenant_id=? AND reason='Web test debit return'",
            (self.tenant_id,)).fetchone()
        after = acct.bill_outstanding(conn, self.tenant_id, bill["id"])
        conn.close()
        self.assertIsNotNone(note)
        self.assertEqual(after, before - note["total_paise"])

        s, notes_page = self.get("/debit-notes")
        self.assertIn("Web test debit return", notes_page)
        s, csv_body = self.get("/export/debit-notes.csv")
        self.assertIn("DBN-", csv_body)
        self.assertIn("Web test debit return", csv_body)

    def test_20_bank_recon_import_match_unmatch(self):
        self.login()
        conn = db.connect(self.db_path)
        bank_line = conn.execute(
            "SELECT l.id AS line_id, l.debit_paise, e.entry_date, e.narration, p.reference"
            " FROM journal_line l JOIN journal_entry e ON e.id=l.entry_id"
            " LEFT JOIN payment p ON p.journal_entry_id=e.id"
            " JOIN account a ON a.id=l.account_id"
            " WHERE l.tenant_id=? AND a.code='1010' AND l.debit_paise>0"
            " AND p.reference IS NOT NULL AND p.reference<>''"
            " ORDER BY e.entry_date DESC LIMIT 1", (self.tenant_id,)).fetchone()
        conn.close()
        self.assertIsNotNone(bank_line, "demo bank receipt not found")

        amount = f'{bank_line["debit_paise"] // 100}.{bank_line["debit_paise"] % 100:02d}'
        csv_text = ("Date,Description,Reference,Debit,Credit\n"
                    f'{bank_line["entry_date"]},NEFT CR DEMO,{bank_line["reference"]},,{amount}\n')
        page = self.get(f"/bank-recon?from={bank_line['entry_date']}&to={bank_line['entry_date']}")[1]
        token = self.csrf(page)
        status, body, final = self.post("/bank-recon/import", {
            "_csrf": token, "account_id": self._bank_account_id(),
            "csv_text": csv_text, "filename": "web-test.csv"})
        self.assertEqual(status, 200)
        self.assertIn("Imported 1 new line", body)
        self.assertIn("bank-recon", final)

        status, body, _ = self.post("/bank-recon/import", {
            "_csrf": token, "account_id": self._bank_account_id(),
            "csv_text": csv_text, "filename": "web-test.csv"})
        self.assertEqual(status, 200)
        self.assertIn("0 new line", body)
        self.assertIn("1 duplicate", body)

        conn = db.connect(self.db_path)
        stmt = conn.execute(
            "SELECT id FROM bank_statement_line WHERE tenant_id=? AND credit_paise=?",
            (self.tenant_id, bank_line["debit_paise"])).fetchone()
        conn.close()
        self.assertIsNotNone(stmt)

        status, body, _ = self.post("/bank-recon/match", {
            "_csrf": token, "statement_id": stmt["id"],
            "journal_line_id": bank_line["line_id"], "account_id": self._bank_account_id(),
            "from": bank_line["entry_date"], "to": bank_line["entry_date"]})
        self.assertEqual(status, 200)
        self.assertIn("Matched", body)
        self.assertIn(bank_line["reference"], body)

        status, body, _ = self.post("/bank-recon/unmatch", {
            "_csrf": token, "statement_id": stmt["id"],
            "account_id": self._bank_account_id(),
            "from": bank_line["entry_date"], "to": bank_line["entry_date"]})
        self.assertEqual(status, 200)
        self.assertIn("Match removed", body)

    def _bank_account_id(self) -> str:
        conn = db.connect(self.db_path)
        row = conn.execute("SELECT id FROM account WHERE tenant_id=? AND code='1010'",
                           (self.tenant_id,)).fetchone()
        conn.close()
        return row["id"]

    def test_21_gstr3b_page_and_csv(self):
        self.login()
        status, page = self.get("/reports/gstr3b?month=2026-06")
        self.assertEqual(status, 200)
        self.assertIn("GSTR-3B", page)
        self.assertIn("3.1 Outward supplies", page)
        self.assertIn("Reconciliation checks", page)
        self.assertIn("6/6 OK", page)
        self.assertIn("Cash payable", page)

        status, csv_body = self.get("/export/gstr3b.csv?month=2026-06")
        self.assertEqual(status, 200)
        self.assertIn("3.1(a)", csv_body)
        self.assertIn("Net eligible ITC", csv_body)
        self.assertIn("Cash payable", csv_body)
        self.assertIn("Reconciliation", csv_body)

        status, bad = self.get("/reports/gstr3b?month=2026-13")
        self.assertEqual(status, 200)
        self.assertIn("YYYY-MM", bad)

    def test_22_team_invite_accept_and_role_change(self):
        self.login()
        status, page = self.get("/users")
        self.assertEqual(status, 200)
        self.assertIn("Invite a team member", page)
        self.assertIn("Owner", page)
        token = self.csrf(page)
        status, created, _ = self.post("/users/invite", {
            "_csrf": token, "email": "web-invite@example.com",
            "full_name": "Web Invite", "role": "accountant"})
        self.assertEqual(status, 200)
        self.assertIn("Invite ready for", created)
        match = re.search(r'value="(http://[^"]+/invite/[A-Za-z0-9_-]+)"', created)
        self.assertIsNotNone(match)
        link = match.group(1)

        with urllib.request.urlopen(link, timeout=15) as r:
            public_page = r.read().decode("utf-8")
        self.assertIn("Join Arthivio Demo Trading", public_page)
        self.assertIn("Create password", public_page)

        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        body = urllib.parse.urlencode({
            "full_name": "Web Invite", "password": "weblongpassword",
            "confirm": "weblongpassword"}).encode()
        req = urllib.request.Request(link, data=body, method="POST")
        with opener.open(req, timeout=15) as r:
            accepted = r.read().decode("utf-8")
            final = r.geturl()
        self.assertTrue(final.endswith("/"))
        self.assertIn("Dashboard", accepted)

        conn = db.connect(self.db_path)
        user_id = conn.execute("SELECT id FROM user WHERE email='web-invite@example.com'").fetchone()["id"]
        conn.close()
        status, page = self.get("/users")
        token = self.csrf(page)
        status, body, _ = self.post("/users/role", {
            "_csrf": token, "user_id": user_id, "role": "manager"})
        self.assertEqual(status, 200)
        self.assertIn("Team updated", body)

    def test_23_shareable_statement_and_revoke(self):
        self.login()
        status, page = self.get(f"/customers/{self.customer_id}")
        self.assertEqual(status, 200)
        self.assertIn("Share and print", page)
        token = self.csrf(page)
        status, page, _ = self.post(f"/customers/{self.customer_id}/share", {
            "_csrf": token, "from": "2026-04-01", "to": "2026-10-31", "days": "7"})
        self.assertEqual(status, 200)
        self.assertIn("Share link created", page)
        match = re.search(r'value="(http://[^"]+/share/statement/[A-Za-z0-9_-]+)"', page)
        self.assertIsNotNone(match)
        link = match.group(1)

        with urllib.request.urlopen(link, timeout=15) as r:
            shared = r.read().decode("utf-8")
            headers = dict(r.headers)
        self.assertIn("Statement of account", shared)
        self.assertIn("Print / Save PDF", shared)
        self.assertIn("noindex", headers.get("X-Robots-Tag", ""))

        conn = db.connect(self.db_path)
        share_id = conn.execute(
            "SELECT id FROM statement_share WHERE tenant_id=? ORDER BY created_at DESC LIMIT 1",
            (self.tenant_id,)).fetchone()["id"]
        conn.close()
        status, body, _ = self.post("/shares/revoke", {
            "_csrf": token, "share_id": share_id,
            "party_type": "customer", "party_id": self.customer_id,
            "from": "2026-04-01", "to": "2026-10-31"})
        self.assertEqual(status, 200)
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(link, timeout=15)
        self.assertEqual(ctx.exception.code, 404)

    def test_24_gstr2b_import_and_export(self):
        self.login()
        status, page = self.get("/gstr2b")
        self.assertEqual(status, 200)
        self.assertIn("GSTR-2B reconciliation", page)
        self.assertIn("Import GSTR-2B working CSV", page)
        csv_text = (
            "GSTIN,Supplier Name,Invoice No,Invoice Date,Taxable Value,IGST,CGST,SGST,Cess,Invoice Value\n"
            "27AAACL1234C1Z5,Loom House Pvt Ltd,LH-2026-097,2026-08-10,42000.00,0,3780.00,3780.00,0,49560.00\n")
        token = self.csrf(page)
        status, body, final = self.post("/gstr2b/import", {
            "_csrf": token, "month": "2026-08", "csv_text": csv_text,
            "filename": "web-gstr2b.csv"})
        self.assertEqual(status, 200)
        self.assertIn("LH-2026-097", body)
        self.assertIn("MATCHED", body)
        self.assertIn("2026-08", final)

        status, csv_body = self.get("/export/gstr2b.csv?month=2026-08")
        self.assertEqual(status, 200)
        self.assertIn("GSTR-2B", csv_body)
        self.assertIn("LH-2026-097", csv_body)
        self.assertIn("Books tax", csv_body)

        status, body, _ = self.post("/gstr2b/import", {
            "_csrf": token, "month": "2026-08", "csv_text": csv_text,
            "filename": "web-gstr2b.csv"})
        self.assertEqual(status, 200)
        self.assertIn("Duplicate", body)
        self.assertNotIn("Imported 0", body)

        conn = db.connect(self.db_path)
        line_id = conn.execute(
            "SELECT id FROM gstr2b_line WHERE tenant_id=? ORDER BY invoice_no LIMIT 1",
            (self.tenant_id,)).fetchone()["id"]
        conn.close()
        status, body, final = self.post("/gstr2b/review", {
            "_csrf": token, "month": "2026-08", "line_id": line_id,
            "status": "itc_hold", "note": "Supplier clarification pending"})
        self.assertEqual(status, 200)
        self.assertIn("Review saved", body)
        self.assertIn("ITC on hold", body)
        self.assertIn("Supplier clarification pending", body)
        self.assertIn("GSTR-3B ITC vs GSTR-2B", body)
        self.assertIn("reviewed=1", final)

        status, csv_body = self.get("/export/gstr2b.csv?month=2026-08")
        self.assertEqual(status, 200)
        self.assertIn("ITC review", csv_body)
        self.assertIn("itc_hold", csv_body)
        self.assertIn("GSTR-3B eligible ITC", csv_body)

        status, body, _ = self.post("/gstr2b/import", {
            "_csrf": token, "month": "2026-08", "csv_text": "not,a,valid\nfile\n",
            "filename": "bad.csv"})
        self.assertEqual(status, 400)
        self.assertIn("header", body.lower())

    def test_25_business_switcher_checks_membership(self):
        self.login()
        conn = db.connect(self.db_path)
        t2 = db.create_tenant(conn, "Second Co", "second", state_code="27")
        owner_id = conn.execute("SELECT id FROM user WHERE email='owner@arthivio.test'").fetchone()["id"]
        security.add_membership(conn, t2, owner_id, "owner")
        conn.close()

        status, page = self.get("/businesses")
        self.assertEqual(status, 200)
        self.assertIn("Arthivio Demo Trading", page)
        self.assertIn("Second Co", page)
        token = self.csrf(page)
        status, body, _ = self.post("/switch", {"_csrf": token, "tenant_id": t2})
        self.assertEqual(status, 200)
        self.assertIn("Second Co", body)
        status, dashboard = self.get("/")
        self.assertIn("Second Co", dashboard)

        status, page = self.get("/businesses")
        token = self.csrf(page)
        status, body, _ = self.post("/switch", {"_csrf": token, "tenant_id": "not-a-membership"})
        self.assertEqual(status, 400)
        self.assertIn("do not have active access", body)

    def test_26_local_demo_autologin_is_opt_in(self):
        old = server.DEMO_AUTOLOGIN
        server.DEMO_AUTOLOGIN = False
        try:
            with self.assertRaises(urllib.error.HTTPError) as ctx:
                self.opener.open(self.url("/demo"), timeout=15)
            self.assertEqual(ctx.exception.code, 404)

            server.DEMO_AUTOLOGIN = True
            jar = http.cookiejar.CookieJar()
            opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
            with opener.open(self.url("/demo"), timeout=15) as r:
                body = r.read().decode("utf-8")
                self.assertEqual(r.status, 200)
                self.assertIn("Dashboard", body)
            conn = db.connect(self.db_path)
            audit = conn.execute(
                "SELECT COUNT(*) FROM audit_log WHERE action='auth.demo_login'").fetchone()[0]
            conn.close()
            self.assertGreaterEqual(audit, 1)
        finally:
            server.DEMO_AUTOLOGIN = old

    def test_27_grouped_dropdown_navigation(self):
        self.login()
        status, body = self.get("/")
        self.assertEqual(status, 200)
        for label in ("Transactions", "Entries", "Books", "Reports", "Masters", "Settings"):
            self.assertIn(f"<span>{label}</span>", body)
        self.assertNotIn("<span>Tax Returns</span>", body)
        self.assertIn('class="menu', body)
        self.assertIn('href="/credit-notes"', body)
        self.assertIn('href="/debit-notes"', body)
        self.assertIn('href="/entries/cash"', body)
        self.assertIn('href="/entries/bank"', body)
        self.assertIn('href="/entries/journal"', body)
        self.assertIn('href="/reports/cash-book"', body)
        self.assertIn('href="/reports/bank-book"', body)
        self.assertIn('href="/reports/analytics"', body)
        self.assertIn('href="/bank-accounts"', body)
        self.assertIn('href="/reports/gstr3b"', body)
        self.assertIn('href="/gstr2b"', body)

    def test_28_ledger_hub_lists_and_creates_accounts(self):
        self.login()
        status, page = self.get("/ledger")
        self.assertEqual(status, 200)
        self.assertIn("Ledger accounts", page)
        self.assertIn("Add ledger account", page)
        self.assertIn("Cash entry", page)
        self.assertIn("Journal entry", page)
        self.assertIn("Rent", page)
        token = self.csrf(page)

        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "6040", "name": "Travel Expense", "type": "expense"})
        self.assertEqual(status, 200)
        self.assertIn("Travel Expense", body)
        self.assertIn("Ledger account created", body)
        self.assertIn("/ledger?account=", final)

        conn = db.connect(self.db_path)
        row = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='6040'",
            (self.tenant_id,)).fetchone()
        audit_count = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='account.create'",
            (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertIsNotNone(row)
        self.assertGreaterEqual(audit_count, 1)

        status, body, _ = self.post("/accounts", {
            "_csrf": token, "code": "6040", "name": "Duplicate Travel", "type": "expense"})
        self.assertEqual(status, 400)
        self.assertIn("already exists", body)

    def test_29_bank_account_details_and_bank_book(self):
        self.login()
        status, page = self.get("/bank-accounts")
        self.assertEqual(status, 200)
        self.assertIn("Add bank account", page)
        self.assertIn("Bank accounts", page)
        token = self.csrf(page)

        status, body, final = self.post("/bank-accounts", {
            "_csrf": token, "code": "1020", "name": "HDFC Current A/c",
            "bank_name": "HDFC Bank", "account_number": "50200099998888",
            "ifsc": "HDFC0001234", "branch": "Andheri West",
            "holder_name": "Arthivio Demo Trading"})
        self.assertEqual(status, 200)
        self.assertIn("HDFC Current A/c", body)
        self.assertIn("Bank account created", body)
        self.assertIn("HDFC Bank", body)
        self.assertIn("/bank-accounts?created=1", final)

        conn = db.connect(self.db_path)
        row = conn.execute(
            "SELECT a.id, b.bank_name, b.account_number, b.ifsc FROM account a"
            " JOIN bank_account b ON b.account_id=a.id"
            " WHERE a.tenant_id=? AND a.code='1020'", (self.tenant_id,)).fetchone()
        conn.close()
        self.assertIsNotNone(row)
        self.assertEqual(row["bank_name"], "HDFC Bank")
        self.assertEqual(row["ifsc"], "HDFC0001234")

        status, book = self.get(f"/reports/bank-book?account={row['id']}")
        self.assertEqual(status, 200)
        self.assertIn("HDFC Current A/c", book)
        self.assertIn("HDFC Bank", book)
        self.assertIn("A/c 50200099998888", book)

        status, recon = self.get(f"/bank-recon?account={row['id']}")
        self.assertEqual(status, 200)
        self.assertIn("HDFC Current A/c", recon)

        status, body, _ = self.post("/bank-accounts", {
            "_csrf": token, "code": "1020", "name": "Duplicate Bank",
            "bank_name": "HDFC Bank", "account_number": "12345678"})
        self.assertEqual(status, 400)
        self.assertIn("already exists", body)

    def test_30_entry_tabs_and_manual_journal_reversal(self):
        self.login()
        status, page = self.get("/entries/cash")
        self.assertEqual(status, 200)
        self.assertIn("Cash entry", page)
        self.assertIn("Row 1 is preselected as Cash", page)
        self.assertIn("Post cash entry", page)

        status, bank_page = self.get("/entries/bank")
        self.assertEqual(status, 200)
        self.assertIn("Bank entry", bank_page)
        self.assertIn("Post bank entry", bank_page)

        conn = db.connect(self.db_path)
        rent = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='6000'",
            (self.tenant_id,)).fetchone()["id"]
        cash = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='1000'",
            (self.tenant_id,)).fetchone()["id"]
        conn.close()
        token = self.csrf(page)
        status, body, final = self.post("/journal", {
            "_csrf": token, "mode": "cash", "entry_date": "2026-10-25",
            "narration": "Reversal test entry",
            "account_id": [rent, cash], "debit": ["500.00", "0"],
            "credit": ["0", "500.00"]})
        self.assertEqual(status, 200)
        self.assertIn("Entry posted", body)
        self.assertIn("/ledger?account=", final)

        conn = db.connect(self.db_path)
        entry = conn.execute(
            "SELECT id, entry_no FROM journal_entry WHERE tenant_id=? AND narration=?",
            (self.tenant_id, "Reversal test entry")).fetchone()
        self.assertIsNotNone(entry)
        conn.close()

        status, reverse_page = self.get(f"/journal?reverse={entry['id']}")
        self.assertEqual(status, 200)
        self.assertIn("Reversal reason", reverse_page)
        self.assertIn(entry["entry_no"], reverse_page)
        reverse_token = self.csrf(reverse_page)
        status, body, final = self.post("/journal/reverse", {
            "_csrf": reverse_token, "entry_id": entry["id"], "reason": "Wrong account"})
        self.assertEqual(status, 200)
        self.assertIn("Reversal entry posted", body)
        self.assertIn("Reversal of", body)
        self.assertIn("reversed=1", final)

        conn = db.connect(self.db_path)
        original = conn.execute(
            "SELECT status FROM journal_entry WHERE id=?", (entry["id"],)).fetchone()
        reversal = conn.execute(
            "SELECT COUNT(*) c FROM journal_entry WHERE tenant_id=? AND source_type='reversal'"
            " AND source_id=?", (self.tenant_id, entry["id"])).fetchone()["c"]
        audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='journal.reverse'",
            (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertEqual(original["status"], "reversed")
        self.assertEqual(reversal, 1)
        self.assertGreaterEqual(audit, 1)

    def test_31_opening_balance_on_creation_and_later(self):
        self.login()
        status, page = self.get("/ledger")
        self.assertEqual(status, 200)
        token = self.csrf(page)

        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "6050", "name": "Insurance", "type": "expense"})
        self.assertEqual(status, 200)
        self.assertIn("Set opening balance", body)
        self.assertIn('action="/opening-balance"', body)
        self.assertIn("Ledger account created", body)

        conn = db.connect(self.db_path)
        account = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='6050'",
            (self.tenant_id,)).fetchone()
        conn.close()
        self.assertIsNotNone(account)
        account_id = account["id"]

        # The same form remains available later when the account page is reopened.
        status, page = self.get(f"/ledger?account={account_id}")
        self.assertEqual(status, 200)
        self.assertIn("Set opening balance", page)
        token = self.csrf(page)
        status, body, final = self.post("/opening-balance", {
            "_csrf": token, "account_id": account_id, "amount": "1,250.50",
            "side": "debit", "as_of": "2026-04-01"})
        self.assertEqual(status, 200)
        self.assertIn("Opening balance posted against Opening Balance Equity", body)
        self.assertIn("Opening balance entry · \u20b91,250.50 Dr", body)
        self.assertIn("opening=1", final)

        conn = db.connect(self.db_path)
        entry = conn.execute(
            "SELECT id, entry_no, source_type, source_id, status FROM journal_entry"
            " WHERE tenant_id=? AND source_type='opening' AND source_id=?",
            (self.tenant_id, account_id)).fetchone()
        lines = conn.execute(
            "SELECT a.code, l.debit_paise, l.credit_paise FROM journal_line l"
            " JOIN account a ON a.id=l.account_id WHERE l.entry_id=? ORDER BY l.line_no",
            (entry["id"],)).fetchall()
        audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='opening_balance.set'",
            (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertEqual(entry["status"], "posted")
        self.assertEqual(lines[0]["code"], "6050")
        self.assertEqual(lines[0]["debit_paise"], 125050)
        self.assertEqual(lines[1]["code"], "3200")
        self.assertEqual(lines[1]["credit_paise"], 125050)
        self.assertGreaterEqual(audit, 1)

        status, body, _ = self.post("/opening-balance", {
            "_csrf": token, "account_id": account_id, "amount": "10.00",
            "side": "debit", "as_of": "2026-04-01"})
        self.assertEqual(status, 400)
        self.assertIn("already set", body)

        status, reverse_page = self.get(f"/journal?reverse={entry['id']}")
        self.assertEqual(status, 200)
        self.assertIn(entry["entry_no"], reverse_page)
        reverse_token = self.csrf(reverse_page)
        status, body, final = self.post("/journal/reverse", {
            "_csrf": reverse_token, "entry_id": entry["id"],
            "reason": "Opening balance corrected"})
        self.assertEqual(status, 200)
        self.assertIn("Reversal entry posted", body)
        self.assertIn("Set opening balance", body)

        token = self.csrf(body)
        status, body, _ = self.post("/opening-balance", {
            "_csrf": token, "account_id": account_id, "amount": "1,300.00",
            "side": "debit", "as_of": "2026-04-01"})
        self.assertEqual(status, 200)
        self.assertIn("Opening balance entry · \u20b91,300.00 Dr", body)

        status, tb = self.get("/reports/trial-balance")
        self.assertIn("Balanced", tb)


    def test_32_full_account_profile_edit_and_party_link(self):
        self.login()
        status, page = self.get("/ledger/account/new")
        self.assertEqual(status, 200)
        for label in ("Account classification", "Sundry Debtors / Accounts Receivable",
                      "Opening balance", "Company, contact and credit details",
                      "GSTIN", "PAN", "Bank details", "Loan / borrowing details",
                      "Credit days", "Industry / business type"):
            self.assertIn(label, page)
        token = self.csrf(page)

        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "1325", "name": "Sunrise Hospital",
            "type": "asset", "group_code": "receivable",
            "legal_name": "Sunrise Hospital Pvt Ltd", "contact_person": "Dr Meera",
            "phone": "9876501234", "email": "accounts@sunrise.test",
            "website": "https://sunrise.test", "address": "Plot 7, Baner Road",
            "city": "Pune", "state_code": "27", "pincode": "411045",
            "country": "India", "gstin": "27AAECS1234F1Z8", "pan": "AAECS1234F",
            "cin": "U85110MH2018PTC123456", "organisation_type": "Private Limited",
            "industry": "Hospital / Healthcare", "credit_days": "45",
            "credit_limit": "250000.00", "notes": "Monthly reconciliation",
            "opening_amount": "5000.00", "opening_side": "debit",
            "opening_as_of": "2026-04-01"})
        self.assertEqual(status, 200)
        for label in ("Account profile", "Sundry Debtors / Accounts Receivable",
                      "Sunrise Hospital Pvt Ltd", "27AAECS1234F1Z8",
                      "Hospital / Healthcare", "45 day(s)", "Open customer statement",
                      "Opening balance entry · ₹5,000.00 Dr"):
            self.assertIn(label, body)
        self.assertIn("created=1", final)

        conn = db.connect(self.db_path)
        account = conn.execute(
            "SELECT * FROM account WHERE tenant_id=? AND code='1325'",
            (self.tenant_id,)).fetchone()
        profile = conn.execute(
            "SELECT * FROM account_profile WHERE tenant_id=? AND account_id=?",
            (self.tenant_id, account["id"])).fetchone()
        customer = conn.execute(
            "SELECT * FROM customer WHERE tenant_id=? AND account_id=?",
            (self.tenant_id, account["id"])).fetchone()
        opening = conn.execute(
            "SELECT * FROM journal_entry WHERE tenant_id=? AND source_type='opening'"
            " AND source_id=?", (self.tenant_id, account["id"])).fetchone()
        conn.close()
        self.assertEqual(account["group_code"], "receivable")
        self.assertEqual(profile["gstin"], "27AAECS1234F1Z8")
        self.assertEqual(profile["credit_days"], 45)
        self.assertEqual(profile["party_id"], customer["id"])
        self.assertIsNotNone(opening)

        status, page = self.get(f"/ledger/account/{account['id']}/edit")
        self.assertEqual(status, 200)
        self.assertIn("Sunrise Hospital Pvt Ltd", page)
        self.assertIn("Save account details", page)
        token = self.csrf(page)
        status, body, final = self.post("/accounts/update", {
            "_csrf": token, "account_id": account["id"],
            "name": "Sunrise Hospital & Research",
            "legal_name": "Sunrise Hospital & Research Pvt Ltd",
            "contact_person": "Dr Meera", "phone": "9999999999",
            "email": "accounts@sunrise.test", "address": "Plot 7, Baner Road",
            "city": "Pune", "state_code": "27", "pincode": "411045",
            "country": "India", "gstin": "27AAECS1234F1Z8", "pan": "AAECS1234F",
            "cin": "U85110MH2018PTC123456", "organisation_type": "Private Limited",
            "industry": "Hospital / Healthcare", "credit_days": "30",
            "credit_limit": "300000.00", "notes": "Revised terms"})
        self.assertEqual(status, 200)
        self.assertIn("Sunrise Hospital", body)
        self.assertIn("Research Pvt Ltd", body)
        self.assertIn("Account details updated", body)
        self.assertIn("30 day(s)", body)
        conn = db.connect(self.db_path)
        audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='account.update'",
            (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertGreaterEqual(audit, 1)

    def test_33_bank_and_loan_accounts_capture_full_details(self):
        self.login()
        status, page = self.get("/ledger/account/new")
        token = self.csrf(page)
        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "1035", "name": "ICICI Current A/c",
            "type": "asset", "group_code": "bank",
            "bank_name": "ICICI Bank", "bank_account_number": "000405001234",
            "ifsc": "ICIC0004050", "branch": "Baner", "holder_name": "Arthivio Demo Trading",
            "opening_amount": "125000.00", "opening_side": "debit",
            "opening_as_of": "2026-04-01"})
        self.assertEqual(status, 200)
        self.assertIn("ICICI Bank", body)
        self.assertIn("A/c 000405001234", body)
        self.assertIn("ICIC0004050", body)
        self.assertIn("Opening balance entry · ₹1,25,000.00 Dr", body)

        status, page = self.get("/ledger/account/new")
        token = self.csrf(page)
        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "2380", "name": "HDFC Term Loan",
            "type": "liability", "group_code": "loans",
            "loan_lender": "HDFC Bank", "loan_account_no": "TL-998877",
            "loan_sanction_date": "2026-04-01", "loan_maturity_date": "2031-03-31",
            "interest_rate": "9.5", "tenure_months": "60", "emi": "12500.50",
            "security_details": "Hypothecation of stock and receivables",
            "opening_amount": "500000.00", "opening_side": "credit",
            "opening_as_of": "2026-04-01"})
        self.assertEqual(status, 200)
        for label in ("HDFC Bank", "TL-998877", "9.50% p.a.", "60 month(s)",
                      "₹12,500.50", "Hypothecation of stock"):
            self.assertIn(label, body)

        conn = db.connect(self.db_path)
        bank = conn.execute(
            "SELECT a.id, b.bank_name, b.account_number, b.ifsc FROM account a"
            " JOIN bank_account b ON b.account_id=a.id"
            " WHERE a.tenant_id=? AND a.code='1035'", (self.tenant_id,)).fetchone()
        self.assertIsNotNone(bank)
        self.assertEqual(bank["ifsc"], "ICIC0004050")
        conn.close()
        status, book = self.get(f"/reports/bank-book?account={bank['id']}")
        self.assertEqual(status, 200)
        self.assertIn("ICICI Current A/c", book)
        self.assertIn("A/c 000405001234", book)

    def test_34_product_master_captures_margin_tax_and_hsn(self):
        self.login()
        status, page = self.get("/items")
        self.assertEqual(status, 200)
        for label in ("Products / services", "Margin % (on cost)", "HSN / SAC",
                      "Sales ledger", "MRP", "Barcode"):
            self.assertIn(label, page)
        token = self.csrf(page)
        status, body, final = self.post("/items", {
            "_csrf": token, "sku": "MED-001", "name": "Surgical Kit",
            "category": "Clinical consumables", "brand": "SteriCare",
            "hsn": "3004", "barcode": "8901234567890", "revenue_code": "4000",
            "cost": "1000.00", "margin": "25", "price": "",
            "mrp": "1500.00", "tax": "12", "stock": "10", "track": "1"})
        self.assertEqual(status, 200)
        self.assertIn("Surgical Kit", body)
        self.assertIn("25.00%", body)
        self.assertIn("12.00%", body)
        self.assertIn("/items?created=1", final)
        conn = db.connect(self.db_path)
        item = conn.execute(
            "SELECT * FROM item WHERE tenant_id=? AND sku='MED-001'",
            (self.tenant_id,)).fetchone()
        conn.close()
        self.assertEqual(item["sale_price_paise"], 125000)
        self.assertEqual(item["margin_bp"], 2500)
        self.assertEqual(item["tax_rate_bp"], 1200)
        self.assertEqual(item["mrp_paise"], 150000)
        self.assertEqual(item["hsn"], "3004")
        self.assertEqual(item["stock_milli"], 10000)

        status, page = self.get("/items")
        token = self.csrf(page)
        status, body, _ = self.post("/items", {
            "_csrf": token, "sku": "MED-002", "name": "Bad MRP",
            "revenue_code": "4000", "cost": "1000.00", "margin": "50",
            "mrp": "1200.00", "tax": "12", "stock": "1", "track": "1"})
        self.assertEqual(status, 400)
        self.assertIn("MRP cannot be lower", body)

    def test_36_bulk_opening_balance_import_preview_post_and_template(self):
        self.login()
        status, ledger = self.get("/ledger")
        self.assertEqual(status, 200)
        token = self.csrf(ledger)

        for code, name, group in (
            ("1605", "Bulk Import Asset", "other_asset"),
            ("2605", "Bulk Import Liability", "other_liability"),
        ):
            status, body, _ = self.post("/accounts", {
                "_csrf": token, "code": code, "name": name, "group_code": group})
            self.assertEqual(status, 200)

        status, page = self.get("/ledger/opening-balances")
        self.assertEqual(status, 200)
        for label in ("Bulk opening balances", "Preview / validate", "Download CSV template",
                      "Opening Balance Equity", "Required columns", "Nothing is written until"):
            self.assertIn(label, page)
        self.assertIn('name="csv_text"', page)
        self.assertIn('name="default_as_of"', page)
        token = self.csrf(page)

        csv_text = (
            "code,side,amount,as_of,note\n"
            "1605,debit,1500.50,2026-04-01,Asset migration\n"
            "2605,credit,500.25,2026-04-01,Liability migration\n"
        )
        status, preview_page, _ = self.post("/ledger/opening-balances", {
            "_csrf": token, "default_as_of": "2026-04-01", "filename": "bulk.csv",
            "csv_text": csv_text, "action": "preview"})
        self.assertEqual(status, 200)
        self.assertIn("2/2 ready", preview_page)
        self.assertIn("Post all 2 opening balances", preview_page)
        self.assertIn("Nothing has been posted yet", preview_page)
        self.assertIn("Asset migration", preview_page)

        conn = db.connect(self.db_path)
        count = conn.execute(
            "SELECT COUNT(*) c FROM journal_entry e JOIN journal_line l ON l.entry_id=e.id"
            " JOIN account a ON a.id=l.account_id WHERE e.tenant_id=? AND"
            " e.source_type='opening' AND a.code IN ('1605','2605')",
            (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertEqual(count, 0)

        status, bad_page, _ = self.post("/ledger/opening-balances", {
            "_csrf": token, "default_as_of": "2026-04-01", "filename": "bad.csv",
            "csv_text": "code,side,amount\n1605,debit,10.00\n9999,debit,5.00\n",
            "action": "preview"})
        self.assertEqual(status, 200)
        self.assertIn("Fix these rows before posting", bad_page)
        self.assertNotIn('name="action" value="post"', bad_page)

        status, done_page, final = self.post("/ledger/opening-balances", {
            "_csrf": token, "default_as_of": "2026-04-01", "filename": "bulk.csv",
            "csv_text": csv_text, "action": "post"})
        self.assertEqual(status, 200)
        self.assertIn("2 opening balance entries posted", done_page)
        self.assertIn("posted=2", final)

        conn = db.connect(self.db_path)
        accounts = {row["code"]: row["id"] for row in conn.execute(
            "SELECT id, code FROM account WHERE tenant_id=? AND code IN ('1605','2605')",
            (self.tenant_id,)).fetchall()}
        openings = conn.execute(
            "SELECT e.id, e.narration FROM journal_entry e WHERE e.tenant_id=?"
            " AND e.source_type='opening' AND e.source_id IN (?,?) ORDER BY e.entry_no",
            (self.tenant_id, accounts["1605"], accounts["2605"])).fetchall()
        audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND"
            " action='opening_balance.bulk_import'", (self.tenant_id,)).fetchone()["c"]
        conn.close()
        self.assertEqual(len(openings), 2)
        self.assertIn("Asset migration", openings[0]["narration"])
        self.assertIn("Liability migration", openings[1]["narration"])
        self.assertGreaterEqual(audit, 1)

        status, tb = self.get("/reports/trial-balance")
        self.assertIn("Balanced", tb)
        status, template = self.get("/ledger/opening-balances/template.csv")
        self.assertIn("code,side,amount,as_of,note", template)
        self.assertIn("25000.00", template)

        # Re-posting the same file is rejected and does not duplicate openings.
        status, duplicate_page, _ = self.post("/ledger/opening-balances", {
            "_csrf": token, "default_as_of": "2026-04-01", "filename": "bulk.csv",
            "csv_text": csv_text, "action": "post"})
        self.assertEqual(status, 400)
        self.assertIn("already set", duplicate_page)
        conn = db.connect(self.db_path)
        count = conn.execute(
            "SELECT COUNT(*) c FROM journal_entry e WHERE e.tenant_id=?"
            " AND e.source_type='opening' AND e.source_id IN (?,?)",
            (self.tenant_id, accounts["1605"], accounts["2605"])).fetchone()["c"]
        conn.close()
        self.assertEqual(count, 2)

    def test_35_account_type_group_selector_covers_business_classifications(self):
        self.login()
        status, page = self.get("/ledger/account/new")
        self.assertEqual(status, 200)
        page_text = html.unescape(page)
        for label in ("Account type / group", "Cash-in-Hand", "Bank Account",
                      "Sundry Debtors / Accounts Receivable",
                      "Sundry Creditors / Accounts Payable", "Inventory / Stock-in-Hand",
                      "Fixed Assets", "Loans (Liability)", "Loans & Advances (Asset)",
                      "Duties & Taxes (GST/TDS)", "Suspense A/c", "Capital Account",
                      "Drawings", "Retained Earnings", "Profit & Loss A/c",
                      "Sales", "Purchase", "Depreciation"):
            self.assertIn(label, page_text)
        self.assertIn('name="group_code"', page)
        self.assertNotIn('name="type"', page)

        token = self.csrf(page)
        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "2395", "name": "Suspense Clearing",
            "group_code": "suspense"})
        self.assertEqual(status, 200)
        self.assertIn("Suspense A/c", html.unescape(body))
        conn = db.connect(self.db_path)
        account = conn.execute(
            "SELECT type, group_code FROM account WHERE tenant_id=? AND code='2395'",
            (self.tenant_id,)).fetchone()
        conn.close()
        self.assertEqual(account["type"], "liability")
        self.assertEqual(account["group_code"], "suspense")
        self.assertIn("/ledger?account=", final)

    def test_37_workspace_offers_four_light_themes_without_navy(self):
        """The signed-in UI has four light palettes and never uses navy app chrome."""
        for token in ("--page:#F6F9FF", "--ink:#212529", "--primary:#2563EB",
                      "--primary-line:#BFDBFE", "--success:#22C55E",
                      '[data-theme="emerald"]', '[data-theme="violet"]',
                      '[data-theme="sunset"]'):
            self.assertIn(token, server.CSS)
        for legacy in ("--ivory", "--emerald", "--sage", "--gold",
                       "#F7F5EF", "#176B4D", "#C6A15B", "#0B1F3A", "#04162F"):
            self.assertNotIn(legacy, server.CSS)
        self.assertEqual(set(server.THEMES), {"ocean", "emerald", "violet", "sunset"})
        for meta in server.THEMES.values():
            self.assertEqual(len(meta["swatches"]), 4)
        self.assertIn("#22D3EE", server.LOGO)
        self.assertIn("var(--primary)", server.PRINT_CSS)
        self.assertNotIn("#2563EB", server.PRINT_CSS)
        self.assertEqual(set(server.LANDING_THEMES), set(server.THEMES))
        for token in ("--lp-hero-a", "--lp-accent", "--lp-primary", "--lp-proof-a",
                      "lp-theme-menu"):
            self.assertIn(token, server.LANDING_CSS)
        self.assertIn("Landing colour", server.landing_page("ocean"))
        violet_landing = server.landing_page("violet")
        self.assertIn('data-theme="violet"', violet_landing)
        self.assertIn("--lp-hero-a:#2D1D52", violet_landing)
        self.assertIn('href="/home?theme=sunset"', violet_landing)
        self.assertIn("Royal Violet", violet_landing)
        self.assertIn('data-theme="ocean"', server.landing_page("navy"))

    def test_38_user_can_choose_a_saved_colour_theme(self):
        self.login()
        status, body = self.get("/settings")
        self.assertEqual(status, 200)
        self.assertIn('action="/preferences/theme"', body)
        for label in ("Ocean Blue", "Emerald Green", "Royal Violet", "Warm Sunset"):
            self.assertIn(label, body)
        token = self.csrf(body)

        status, body, final = self.post(
            "/preferences/theme", {"_csrf": token, "theme": "violet"})
        self.assertEqual(status, 200)
        self.assertIn("theme_saved=1", final)
        self.assertIn('data-theme="violet"', body)
        self.assertIn("Theme saved", body)
        self.assertIn("Royal Violet", body)

        conn = db.connect(self.db_path)
        user = conn.execute(
            "SELECT theme FROM user WHERE email='owner@arthivio.test'").fetchone()
        audit = conn.execute(
            "SELECT meta FROM audit_log WHERE action='user.theme_updated'"
            " ORDER BY created_at DESC LIMIT 1").fetchone()
        conn.close()
        self.assertEqual(user["theme"], "violet")
        self.assertIn('"to": "violet"', audit["meta"])

        status, dashboard = self.get("/")
        self.assertIn('data-theme="violet"', dashboard)
        status, login = self.get("/login")
        self.assertIn('data-theme="violet"', login)

        conn = db.connect(self.db_path)
        invoice_id = conn.execute(
            "SELECT id FROM invoice WHERE tenant_id=? LIMIT 1", (self.tenant_id,)).fetchone()["id"]
        conn.close()
        status, printed = self.get(f"/invoices/{invoice_id}/print")
        self.assertEqual(status, 200)
        self.assertIn('data-theme="violet"', printed)
        self.assertIn("--primary:#7C3AED", printed)

        status, settings = self.get("/settings")
        token = self.csrf(settings)
        status, rejected, _ = self.post(
            "/preferences/theme", {"_csrf": token, "theme": "navy"})
        self.assertEqual(status, 400)
        self.assertIn("Choose a valid colour theme", rejected)

        status, landing = self.get("/home?theme=emerald")
        self.assertIn('data-theme="emerald"', landing)
        self.assertIn("--lp-hero-a:#0D3F31", landing)
        status, landing = self.get("/home")
        self.assertIn('data-theme="emerald"', landing)
        status, landing = self.get("/home?theme=ocean")
        self.assertIn('data-theme="ocean"', landing)

        status, settings = self.get("/settings")
        token = self.csrf(settings)
        status, _, _ = self.post(
            "/preferences/theme", {"_csrf": token, "theme": "ocean"})
        self.assertEqual(status, 200)

    def test_39_account_archive_merge_and_restore(self):
        self.login()
        status, page = self.get("/ledger/account/new")
        self.assertEqual(status, 200)
        token = self.csrf(page)

        for code, name, amount in (
            ("6081", "Old Office Expenses", "1000.00"),
            ("6082", "Office Expenses", "500.00"),
        ):
            status, body, final = self.post("/accounts", {
                "_csrf": token, "code": code, "name": name,
                "group_code": "other_expense",
                "opening_amount": amount, "opening_side": "debit",
                "opening_as_of": "2026-04-01"})
            self.assertEqual(status, 200)
            self.assertIn("created=1", final)

        conn = db.connect(self.db_path)
        ids = {row["code"]: row["id"] for row in conn.execute(
            "SELECT id, code FROM account WHERE tenant_id=? AND code IN ('6081','6082')",
            (self.tenant_id,)).fetchall()}
        cash_id = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='1000'",
            (self.tenant_id,)).fetchone()["id"]
        conn.close()
        source_id, target_id = ids["6081"], ids["6082"]

        status, page = self.get(f"/ledger/account/{source_id}/archive")
        self.assertEqual(status, 200)
        for label in ("Archive or merge", "Archive is available only at an all-time balance of zero",
                      "Merge into another ledger", "Merge reason", "All-time balance"):
            self.assertIn(label, page)
        self.assertNotIn('name="archive_reason"', page)
        self.assertIn(f'value="{target_id}"', page)
        token = self.csrf(page)

        status, body, final = self.post("/accounts/merge", {
            "_csrf": token, "account_id": source_id, "target_id": target_id,
            "reason": "Duplicate office-expense ledger"})
        self.assertEqual(status, 200)
        self.assertIn("merged=1", final)
        self.assertIn("source ledger was merged into this account", body)
        self.assertIn("Archived (", body)

        conn = db.connect(self.db_path)
        source = conn.execute("SELECT * FROM account WHERE id=?", (source_id,)).fetchone()
        target = conn.execute("SELECT * FROM account WHERE id=?", (target_id,)).fetchone()
        merge_entry = conn.execute(
            "SELECT * FROM journal_entry WHERE tenant_id=? AND source_type='merge'"
            " AND source_id=?", (self.tenant_id, source_id)).fetchone()
        merge_audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='account.merge'"
            " AND target_id=?", (self.tenant_id, source_id)).fetchone()["c"]
        conn.close()
        self.assertEqual(source["status"], "archived")
        self.assertEqual(source["merged_into_id"], target_id)
        self.assertEqual(target["status"], "active")
        self.assertIsNotNone(merge_entry)
        self.assertGreaterEqual(merge_audit, 1)

        status, ledger = self.get(f"/ledger?account={source_id}")
        self.assertEqual(status, 200)
        self.assertIn("Archived ledger", ledger)
        self.assertIn("Merged into", ledger)
        self.assertIn("6082", ledger)

        status, page = self.get("/ledger/account/new")
        token = self.csrf(page)
        status, body, final = self.post("/accounts", {
            "_csrf": token, "code": "6083", "name": "Archive Me",
            "group_code": "other_expense"})
        self.assertEqual(status, 200)
        self.assertIn("created=1", final)
        conn = db.connect(self.db_path)
        archive_id = conn.execute(
            "SELECT id FROM account WHERE tenant_id=? AND code='6083'",
            (self.tenant_id,)).fetchone()["id"]
        conn.close()

        status, page = self.get(f"/ledger/account/{archive_id}/archive")
        token = self.csrf(page)
        status, body, final = self.post("/accounts/archive", {
            "_csrf": token, "account_id": archive_id, "reason": "Obsolete test ledger"})
        self.assertEqual(status, 200)
        self.assertIn("archived=1", final)
        self.assertIn("Account archived", body)
        conn = db.connect(self.db_path)
        archived = conn.execute(
            "SELECT status, archived_at FROM account WHERE id=?", (archive_id,)).fetchone()
        archive_audit = conn.execute(
            "SELECT COUNT(*) c FROM audit_log WHERE tenant_id=? AND action='account.archive'"
            " AND target_id=?", (self.tenant_id, archive_id)).fetchone()["c"]
        conn.close()
        self.assertEqual(archived["status"], "archived")
        self.assertIsNotNone(archived["archived_at"])
        self.assertGreaterEqual(archive_audit, 1)

        status, journal = self.get("/entries/journal")
        token = self.csrf(journal)
        status, body, _ = self.post("/journal", {
            "_csrf": token, "mode": "journal",
            "narration": "Blocked archived posting",
            "account_id": [archive_id, cash_id], "debit": ["10.00", ""],
            "credit": ["", "10.00"]})
        self.assertEqual(status, 400)
        self.assertIn("archived", body.lower())

        status, page = self.get(f"/ledger/account/{archive_id}/archive")
        self.assertIn("Restore this ledger", page)
        token = self.csrf(page)
        status, body, final = self.post("/accounts/restore", {
            "_csrf": token, "account_id": archive_id, "reason": "Archived by mistake"})
        self.assertEqual(status, 200)
        self.assertIn("restored=1", final)
        self.assertIn("Account restored", body)
        conn = db.connect(self.db_path)
        restored = conn.execute(
            "SELECT status, merged_into_id FROM account WHERE id=?", (archive_id,)).fetchone()
        conn.close()
        self.assertEqual(restored["status"], "active")
        self.assertIsNone(restored["merged_into_id"])

        status, page = self.get(f"/ledger/account/{cash_id}/archive")
        self.assertEqual(status, 200)
        self.assertIn("Archive and merge are locked", page)

    def test_40_analytics_api_contract_and_auth(self):
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            opener.open(self.url("/api/v1/analytics/overview"), timeout=15)
        self.assertEqual(ctx.exception.code, 401)
        try:
            unauth = json.loads(ctx.exception.read().decode("utf-8"))
        finally:
            ctx.exception.close()
        self.assertEqual(unauth["code"], "UNAUTHENTICATED")

        self.login()
        status, catalog = self.get_json("/api/v1/analytics/catalog")
        self.assertEqual(status, 200)
        self.assertEqual(catalog["schema_version"], "1.0")
        self.assertEqual(catalog["permission"], "report.view")
        keys = {m["key"] for m in catalog["metrics"]}
        self.assertIn("revenue_paise", keys)
        self.assertIn("trial_balance_balanced", keys)
        self.assertTrue(all(m["kind"] == "actual" for m in catalog["metrics"]))

        status, data = self.get_json("/api/v1/analytics/overview")
        self.assertEqual(status, 200)
        self.assertEqual(data["schema_version"], "1.0")
        self.assertIn("from", data["period"])
        self.assertIn("to", data["period"])
        self.assertIn("revenue_paise", data["metrics"])
        self.assertIn("monthly", data)
        self.assertGreaterEqual(len(data["monthly"]["months"]), 1)
        self.assertEqual(data["checks"]["unbalanced_entries"], 0)
        self.assertTrue(data["metrics"]["trial_balance_balanced"])
        self.assertGreaterEqual(data["freshness"]["journal_entries"], 1)
        self.assertIn("expense_breakdown", data)
        self.assertIn("top_customers", data)

        status, bad = self.get_json("/api/v1/analytics/overview?from=not-a-date")
        self.assertEqual(status, 400)
        self.assertEqual(bad["code"], "BAD_DATE")

        status, page = self.get("/reports/analytics")
        self.assertEqual(status, 200)
        self.assertIn("Analytics overview", page)
        self.assertIn("Revenue and expenses by month", page)
        self.assertIn("Actuals only", page)
        self.assertIn("Top customers", page)
        self.assertNotIn("NaN", page)
        self.assertNotIn("None", page)

    def test_41_excel_workbook_download(self):
        self.login()
        status, page = self.get("/reports/analytics")
        self.assertEqual(status, 200)
        self.assertIn("Download Excel workbook", page)

        status, headers, data = self.get_bytes("/export/workbook.xlsx")
        self.assertEqual(status, 200)
        self.assertEqual(
            headers.get("Content-Type"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", headers.get("Content-Disposition", ""))
        self.assertEqual(headers.get("Cache-Control"), "no-store")
        self.assertTrue(data.startswith(b"PK"))
        wb = load_workbook(io.BytesIO(data), data_only=False)
        self.assertIn("Executive KPIs", wb.sheetnames)
        self.assertIn("AI Insights", wb.sheetnames)
        self.assertIn("Monthly", wb.sheetnames)
        self.assertGreaterEqual(len(wb["Monthly"]._charts), 2)

        try:
            self.opener.open(self.url("/export/workbook.xlsx?from=not-a-date"), timeout=15)
        except urllib.error.HTTPError as e:
            try:
                self.assertEqual(e.code, 400)
                self.assertIn("Invalid period", e.read().decode("utf-8"))
            finally:
                e.close()
        else:
            self.fail("Expected an invalid-date workbook request to fail.")

    def test_42_powerbi_pack_download(self):
        self.login()
        status, page = self.get("/reports/analytics")
        self.assertEqual(status, 200)
        self.assertIn("Power BI pack", page)

        status, headers, data = self.get_bytes(
            "/export/powerbi.zip?from=2026-04-01&to=2026-10-31")
        self.assertEqual(status, 200)
        self.assertEqual(headers.get("Content-Type"), "application/zip")
        self.assertIn(".zip", headers.get("Content-Disposition", ""))
        self.assertEqual(headers.get("Cache-Control"), "no-store")
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            self.assertIn("manifest.json", zf.namelist())
            self.assertIn("fact_journal.csv", zf.namelist())
            manifest = json.loads(zf.read("manifest.json"))
            self.assertTrue(manifest["actuals_only"])
            self.assertEqual(manifest["period"]["from"], "2026-04-01")
            self.assertGreaterEqual(manifest["journal_rows_exported"], 1)

        try:
            self.opener.open(self.url("/export/powerbi.zip?from=bad"), timeout=15)
        except urllib.error.HTTPError as e:
            try:
                self.assertEqual(e.code, 400)
                self.assertIn("Invalid period", e.read().decode("utf-8"))
            finally:
                e.close()
        else:
            self.fail("Expected an invalid-date Power BI request to fail.")

    def test_43_echarts_static_asset_and_interactive_view(self):
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        with opener.open(self.url("/static/echarts.min.js"), timeout=30) as r:
            body = r.read()
            self.assertEqual(r.status, 200)
            self.assertIn("application/javascript", r.headers.get("Content-Type", ""))
            self.assertGreater(len(body), 500000)
            self.assertIn(b"apache", body[:500].lower())
        self.login()
        status, page = self.get("/reports/analytics")
        self.assertEqual(status, 200)
        self.assertIn('id="echarts-monthly"', page)
        self.assertIn('id="echarts-expense"', page)
        self.assertIn("/static/echarts.min.js", page)
        self.assertIn("Interactive monthly view", page)
        self.assertIn("Expense breakdown", page)
        self.assertIn("ECharts", page)

    def test_44_grounded_insight_api_and_page(self):
        jar = http.cookiejar.CookieJar()
        opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            opener.open(self.url("/api/v1/analytics/insights"), timeout=15)
        self.assertEqual(ctx.exception.code, 401)
        ctx.exception.close()

        self.login()
        status, payload = self.get_json("/api/v1/analytics/insights")
        self.assertEqual(status, 200)
        self.assertEqual(payload["engine"], "deterministic_rules_v1")
        self.assertFalse(payload["generative_ai"])
        self.assertFalse(payload["forecast_included"])
        self.assertTrue(payload["actuals_only"])
        keys = {o["key"] for o in payload["observations"]}
        self.assertIn("net_profit_margin", keys)
        self.assertIn("journal_checks", keys)
        self.assertIn("basis", payload["observations"][0])
        self.assertIn("no external model was called", payload["summary"].lower())

        status, bad = self.get_json("/api/v1/analytics/insights?from=bad")
        self.assertEqual(status, 400)
        self.assertEqual(bad["code"], "BAD_DATE")

        status, page = self.get("/reports/analytics")
        self.assertEqual(status, 200)
        self.assertIn("AI Insights", page)
        self.assertIn("Grounded rules", page)
        self.assertIn("Open insight JSON", page)
        self.assertIn("No external model was called", page)

        status, page = self.get("/reports/analytics?from=not-a-date")
        self.assertEqual(status, 200)
        self.assertIn("Could not use that period", page)
        self.assertIn("Analytics overview", page)


    def test_45_real_xlsx_multipart_imports(self):
        self.login()
        conn = db.connect(self.db_path)
        try:
            account = conn.execute(
                """SELECT a.code FROM account a
                   WHERE a.tenant_id=? AND a.is_postable=1 AND a.status='active'
                     AND NOT EXISTS (
                       SELECT 1 FROM journal_line l
                       WHERE l.account_id=a.id AND l.tenant_id=a.tenant_id
                     )
                   ORDER BY a.code LIMIT 1""",
                (self.tenant_id,),
            ).fetchone()
            self.assertIsNotNone(account, "demo needs one unused postable account")
            bank_row = conn.execute(
                "SELECT id FROM account WHERE tenant_id=? AND group_code='bank' LIMIT 1",
                (self.tenant_id,),
            ).fetchone()
        finally:
            conn.close()

        status, page = self.get("/ledger/opening-balances")
        self.assertEqual(status, 200)
        csrf = self.csrf(page)
        data = self.xlsx_bytes([
            ["code", "side", "amount", "as_of", "note"],
            [account["code"], "debit", 1234.50, "2026-04-01", "XLSX import test"],
        ])
        status, body, _ = self.post_multipart(
            "/ledger/opening-balances",
            {"_csrf": csrf, "default_as_of": "2026-04-01", "action": "preview"},
            {"xlsx_file": ("opening.xlsx", data,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(status, 200)
        self.assertIn("1/1 ready", body)
        self.assertIn("XLSX import test", body)

        if bank_row is None:
            self.skipTest("demo has no bank account for the statement import check")
        status, page = self.get(f"/bank-recon?account={bank_row['id']}")
        self.assertEqual(status, 200)
        csrf = self.csrf(page)
        data = self.xlsx_bytes([
            ["Date", "Description", "Reference", "Debit", "Credit", "Balance"],
            ["2026-06-15", "XLSX NEFT TEST", "XLSX-1", None, 4321.00, 54321.00],
        ])
        status, body, url = self.post_multipart(
            "/bank-recon/import",
            {"_csrf": csrf, "account_id": bank_row["id"]},
            {"xlsx_file": ("statement.xlsx", data,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(status, 200)
        self.assertIn("imported=1", url)
        self.assertIn("Bank reconciliation", body)


if __name__ == "__main__":
    unittest.main(verbosity=2)
