"""Real .xlsx to CSV conversion tests for the existing import validators."""
from __future__ import annotations

import io
import os
import sys
import unittest
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook

from app import bank, db, gstr2b, opening_import, xlsx_import


def workbook_bytes(rows, title="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class XlsxImportTest(unittest.TestCase):
    def setUp(self):
        self.conn = db.connect(":memory:")
        db.init_db(self.conn)
        self.t = db.create_tenant(self.conn, "XLSX Co", "xlsx-co", state_code="27")

    def tearDown(self):
        self.conn.close()

    def test_opening_balance_xlsx_becomes_valid_preview(self):
        data = workbook_bytes([
            ["code", "side", "amount", "as_of", "note"],
            ["1000", "debit", 25000.00, date(2026, 4, 1), "Cash opening"],
            ["2000", "credit", "18000.50", "2026-04-01", "Supplier opening"],
        ])
        text = xlsx_import.xlsx_to_csv(data)
        self.assertIn("1000,debit,25000,2026-04-01,Cash opening", text)
        self.assertIn("2000,credit,18000.50,2026-04-01,Supplier opening", text)
        preview = opening_import.preview(self.conn, self.t, text, "2026-04-01")
        self.assertEqual(preview["count"], 2)
        self.assertEqual(preview["valid_count"], 2)
        self.assertEqual(preview["errors"], [])
        self.assertEqual(preview["total_debit"], 2500000)
        self.assertEqual(preview["total_credit"], 1800050)

    def test_bank_statement_xlsx_parses_dates_and_paise(self):
        data = workbook_bytes([
            ["Date", "Description", "Reference", "Debit", "Credit", "Balance"],
            [date(2026, 6, 1), "NEFT CR SUNDARAM", "NEFT-88421", None, 23600.00, 523600.00],
            [date(2026, 6, 2), "BANK CHARGES", "CHG-1", 125.50, None, 523474.50],
        ])
        text = xlsx_import.xlsx_to_csv(data)
        rows = bank.parse_csv(text)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["txn_date"], "2026-06-01")
        self.assertEqual(rows[0]["credit_paise"], 2360000)
        self.assertEqual(rows[1]["debit_paise"], 12550)

    def test_gstr2b_xlsx_parses_dates_and_tax(self):
        data = workbook_bytes([
            ["Supplier GSTIN", "Supplier Name", "Invoice No", "Invoice Date",
             "Taxable Value", "IGST", "CGST", "SGST", "Cess", "Invoice Value"],
            ["27AAACL1234C1Z5", "Loom House Pvt Ltd", "LH-2026-097", date(2026, 8, 10),
             42000.00, 0, 3780.00, 3780.00, 0, 49560.00],
        ])
        text = xlsx_import.xlsx_to_csv(data)
        rows = gstr2b.parse_csv(text)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["invoice_date"], "2026-08-10")
        self.assertEqual(rows[0]["taxable_paise"], 4200000)
        self.assertEqual(rows[0]["tax_paise"], 756000)
        self.assertEqual(rows[0]["total_paise"], 4956000)

    def test_sheet_selection_header_and_formula_boundary(self):
        wb = Workbook()
        first = wb.active
        first.title = "Read me"
        first["A1"] = "This sheet is not the data"
        second = wb.create_sheet("Data")
        second.append(["Code", "Side", "Amount"])
        second.append(["1000", "debit", 10.00])
        second.append(["=SUM(C2:C2)", "credit", 5.00])
        buf = io.BytesIO()
        wb.save(buf)
        text = xlsx_import.xlsx_to_csv(buf.getvalue(), "Data")
        rows = list(__import__("csv").reader(io.StringIO(text)))
        self.assertEqual(rows[0], ["Code", "Side", "Amount"])
        self.assertEqual(rows[1], ["1000", "debit", "10"])
        self.assertEqual(rows[2][0], "")
        self.assertEqual(rows[2][1], "credit")

    def test_invalid_files_are_rejected(self):
        with self.assertRaises(xlsx_import.XlsxImportError):
            xlsx_import.xlsx_to_csv(b"")
        with self.assertRaises(xlsx_import.XlsxImportError):
            xlsx_import.xlsx_to_csv(b"not an xlsx")
        with self.assertRaises(xlsx_import.XlsxImportError):
            xlsx_import.xlsx_to_csv(workbook_bytes([[]]))
        with self.assertRaises(xlsx_import.XlsxImportError):
            xlsx_import.xlsx_to_csv(workbook_bytes([["code"]]), "Missing")


if __name__ == "__main__":
    unittest.main()
